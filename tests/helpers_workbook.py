"""Build minimal in-memory Direct Shades–style workbooks for tests."""

from __future__ import annotations

import io

from openpyxl import Workbook


def build_window_matrix_xlsx(
    *,
    total_shades: int = 5,
    markings: list[tuple[str, int]] | None = None,
    blind_rows: list[dict] | None = None,
    blind_header: list[str] | None = None,
) -> bytes:
    markings = markings or [("A-101", 2), ("B-202", 3)]
    wb = Workbook()
    ws = wb.active
    ws.title = "WINDOW MATRIX"
    ws.append(["Project Name", "Test Tower"])
    ws.append([])
    ws.append(["WINDOW MARKINGS", "SHADES PER OPENING", "TOTAL SHADES", "TOTAL WINDOWS"])
    for mark, qty in markings:
        ws.append([mark, 1, qty, qty])
    ws.append(["TOTAL", "", total_shades, total_shades])

    bq = wb.create_sheet("Blind QTY UNITS LIVINGS")
    bq.append(["Area", "LIVINGS"])
    bq.append(["Unit Quantity", 10])
    header = blind_header or ["Window Tag", "QTY", "W", "H", "SYSTEM", "ROOM"]
    bq.append(header)
    rows = blind_rows or [
        {"tag": "A-101", "qty": 1, "w": 48, "h": 72, "sys": "Roller", "room": "Living"},
        {"tag": "B-202", "qty": 1, "w": 36, "h": 60, "sys": "Blackout", "room": "Bed"},
    ]
    for r in rows:
        bq.append([r["tag"], r["qty"], r["w"], r["h"], r.get("sys"), r.get("room")])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_multi_section_matrix_xlsx(
    *,
    total_shades: int = 6,
    metadata_rows: int = 6,
) -> bytes:
    """
    Window Matrix whose Blind QTY sheet stacks several areas.

    Exercises the harder real-world shape: a metadata block above the header, area
    label rows between data blocks, sub-total rows mid-sheet, and the same window tag
    reused in two areas.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "WINDOW MATRIX"
    ws.append(["Project Name", "Madeira Tower"])
    ws.append([])
    ws.append(["WINDOW MARKINGS", "SHADES PER OPENING", "TOTAL SHADES", "TOTAL WINDOWS"])
    ws.append(["LIVING AREAS", "", "", ""])
    ws.append(["A-101", 1, 3, 3])
    ws.append(["BEDROOMS", "", "", ""])
    ws.append(["A-101", 1, 3, 3])
    ws.append(["TOTAL", "", total_shades, total_shades])

    bq = wb.create_sheet("Blind QTY UNITS")
    bq.append(["Project", "Madeira Tower"])
    bq.append(["Area", "UNITS"])
    bq.append(["Unit Quantity", 12])
    for _ in range(max(0, metadata_rows - 3)):
        bq.append([])
    bq.append(["Window Tag", "QTY", 'WIDTH (IN)', "HEIGHT (IN)", "SYSTEM", "ROOM"])
    bq.append(["LIVINGS", None, None, None, None, None])
    bq.append(["A-101", 2, 48, 72, "Solar Roller", "Living"])
    bq.append(["A-102", 1, 36, 60, "Blackout", "Living"])
    bq.append(["SUBTOTAL", 3, None, None, None, None])
    bq.append(["BEDROOMS", None, None, None, None, None])
    bq.append(["A-101", 2, 30, 66, "Blackout Motorized", "Bed 1"])
    bq.append(["A-201", 1, 40, 80, "Dual", "Bed 2"])
    bq.append(["TOTAL", 6, None, None, None, None])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_bid_summary_xlsx(
    *,
    tabs: dict[str, list[dict]] | None = None,
    include_total_row: bool = True,
) -> bytes:
    """Bid Summary workbook with one priced tab per unit type."""
    tabs = tabs or {
        "UNITS": [
            {"description": "Solar roller shades — units", "qty": 4, "unit": 300.0},
            {"description": "Blackout shades — units", "qty": 2, "unit": 400.0},
        ],
        "COM-4\"": [
            {"description": "Common area 4\" fascia shades", "qty": 2, "unit": 500.0},
        ],
    }

    wb = Workbook()
    first = True
    for sheet_name, lines in tabs.items():
        ws = wb.active if first else wb.create_sheet(sheet_name)
        ws.title = sheet_name
        first = False
        ws.append(["Project Name", "Madeira Tower"])
        ws.append([])
        ws.append(["DESCRIPTION", "QTY", "UNIT PRICE", "AMOUNT"])
        total = 0.0
        for line in lines:
            amount = line["qty"] * line["unit"]
            total += amount
            ws.append([line["description"], line["qty"], line["unit"], amount])
        if include_total_row:
            ws.append(["TOTAL", "", "", total])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_material_summary_xlsx() -> bytes:
    """Material Summary workbook with fabric / system lines per tab."""
    wb = Workbook()
    ws = wb.active
    ws.title = "UNITS"
    ws.append(["Project Name", "Madeira Tower"])
    ws.append(["Fabric", "Phifer SheerWeave 5% Charcoal"])
    ws.append(["System", "4\" fascia roller, 1.5\" tube"])
    ws.append(["Motor", "Somfy Sonesse 30 RTS"])

    blk = wb.create_sheet("BLK")
    blk.append(["Fabric", "Blackout vinyl, white reverse"])
    blk.append(["System", "Cassette with side channels"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
