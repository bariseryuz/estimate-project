"""A project that does not use Madeira's file names, sheet names, or area names."""

import io
import unittest

from openpyxl import Workbook

from domain.direct_shades_workbooks import WorkbookKind, merge_workbook_analyses, parse_workbook_xlsx
from domain.pricing import build_estimate_from_takeoff, price_takeoff
from domain.workbook_takeoff import build_takeoff_from_workbook
from pipeline.project_summary import build_project_summary
from pipeline.quantity_schedule import build_quantity_schedule
from pipeline.validation_adjust import adjust_validation_for_workbook


def _harbor() -> bytes:
    """
    Same column keys as a Direct Shades set, none of the Madeira labels.

    The matrix tab is 'Schedule', quantity tabs are 'Tower Openings' and
    'Pool Openings', and the price tabs sit in the same file.
    """
    wb = Workbook()
    schedule = wb.active
    schedule.title = "Schedule"
    schedule.append(["Job", "", "", "Project Name", "Harbor Tower"])
    schedule.append(["WINDOW MARKINGS", "SHADES PER OPENING", "TOTAL WINDOWS", "TOTAL SHADES"])
    schedule.append(["TOWER A"])
    schedule.append(["A", 1, 4, 4])
    schedule.append(["GRAND TOTAL", "", 4, 4])

    tower = wb.create_sheet("Tower Openings")
    tower.append(["Area", "TOWER A"])
    tower.append(["Window Tag", "QTY", "W (INCH)", "H (INCH)", "SYSTEM"])
    tower.append(["A", 4, 41, 96, "Fascia"])

    pool = wb.create_sheet("Pool Openings")
    pool.append(["Area", "POOL HOUSE"])
    pool.append(["Mark", "Count", "Width (in)", "Height (in)", "Type"])
    pool.append(["C", 2, 36, 72, "Roller"])

    price = wb.create_sheet("Price List")
    price.append(["Window Marking", "W", "H", "Qty", "Sales Price", "Total Price"])
    price.append(["TOWER A"])
    price.append(["A", 41, 108, 4, 50, 200])
    price.append(["POOL HOUSE"])
    price.append(["C", 36, 84, 2, 40, 80])
    price.append(["Total Product", "", "", "", "", 280])
    price.append(["", "", "", ""])
    price.append(["", "Tax", 0.07, 19.60])
    price.append(["", "Total Bid", "", 299.60])

    dark = wb.create_sheet("Dark Fabric")
    dark.append(["Description", "Qty", "Unit Price", "Amount"])
    dark.append(["A", 4, 90, 360])
    dark.append(["Total", "", "", 360])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestDocumentKeys(unittest.TestCase):
    def test_headers_identify_the_workbook_without_madeira_names(self):
        parsed = parse_workbook_xlsx("Harbor Takeoff.xlsx", _harbor())
        self.assertEqual(parsed.kind, WorkbookKind.WINDOW_MATRIX)
        self.assertEqual(parsed.project_name, "Harbor Tower")
        self.assertEqual(parsed.matrix_total_shades, 4)
        sections = {row["section"] for row in parsed.matrix_rows}
        self.assertEqual(sections, {"TOWER A"})
        self.assertEqual(parsed.matrix_rows[0]["marking"], "A")

        tags = {(line["windowTag"], line["section"], line["quantity"]) for line in parsed.blind_qty_lines}
        self.assertIn(("A", "TOWER A", 4), tags)
        self.assertIn(("C", "POOL HOUSE", 2), tags)

        merged = merge_workbook_analyses([parsed])
        self.assertEqual(merged["projectShadeCount"], 6)
        note = merged["measurementReference"]["note"]
        self.assertIn("W (INCH)", note)
        self.assertIn("Width (in)", note)
        self.assertIn("12 in greater", note)
        self.assertIn("opening", note.lower())

        self.assertEqual(merged["additionalShadeCount"], 2)
        self.assertEqual(merged["referencePricing"]["grandTotal"], 280)
        self.assertEqual(merged["referencePricing"]["clientTotal"], 299.60)
        self.assertIn("Dark Fabric", merged["referencePricing"]["alternateSheets"])
        self.assertNotIn("Dark Fabric", merged["referencePricing"]["quotedSheets"])

        takeoff = build_takeoff_from_workbook(merged)
        self.assertEqual(takeoff["totalShadeCount"], 6)
        priced = price_takeoff(takeoff["takeoffItems"], workbook=merged)
        self.assertEqual(priced["subtotal"], 280)
        self.assertEqual(priced["total"], 299.60)

        estimate = build_estimate_from_takeoff(takeoff, merged)
        validation = adjust_validation_for_workbook(
            {"confidenceScore": 70}, merged, takeoff, estimate
        )
        count_check = next(
            c for c in validation["crossChecks"]
            if c["item"] == "Total shade count vs WINDOW MATRIX TOTAL row"
        )
        self.assertTrue(count_check["match"])
        self.assertIn("4 on the matrix + 2", count_check["documentValue"])
        price_check = next(
            c for c in validation["crossChecks"]
            if c["item"] == "Project total vs Bid Summary grand total"
        )
        self.assertTrue(price_check["match"])
        self.assertIn("299.60", price_check["documentValue"])

        schedule = build_quantity_schedule(takeoff, estimate, merged, {})
        summary = build_project_summary(
            source_meta={"fileNames": ["Harbor Takeoff.xlsx"]},
            workbook=merged,
            takeoff=takeoff,
            estimation=estimate,
            validation=validation,
            quantity_schedule=schedule,
        )
        self.assertEqual(summary["matchStatus"], "match")
        self.assertIn("4 on the window matrix + 2", summary["matchNote"])
        self.assertEqual(schedule["methodology"]["measurementReference"]["note"], note)

        from pipeline.client_offer import build_client_offer_package
        from pipeline.estimation_audit import build_estimation_audit

        offer = build_client_offer_package(estimate, takeoff, merged, validation, {}, schedule)
        self.assertIn("W (INCH)", offer["measurementNote"])
        self.assertIn(offer["measurementNote"], offer["emailDraft"])
        audit = build_estimation_audit(
            source_meta={"fileNames": ["Harbor Takeoff.xlsx"]},
            workbook=merged,
            takeoff=takeoff,
            estimation=estimate,
            context={},
            quantity_schedule=schedule,
        )
        measure_ref = next(r for r in audit["documentReferences"] if r["role"] == "Measurement reference")
        self.assertIn("Tower Openings: W (INCH)", measure_ref["location"])
        self.assertIn("12 in greater", measure_ref["usedFor"])


if __name__ == "__main__":
    unittest.main()
