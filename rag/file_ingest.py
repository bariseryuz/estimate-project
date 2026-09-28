"""
Extract text (and optional PDF/image payloads) from uploaded project files.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pypdf

# Extensions we accept in the UI / API (lowercase, with dot)
SUPPORTED_EXTENSIONS = frozenset(
    {
        ".pdf",
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".gif",
        ".bmp",
        ".tif",
        ".tiff",
        ".xlsx",
        ".xlsm",
        ".xls",
        ".csv",
        ".tsv",
        ".docx",
        ".pptx",
        ".txt",
        ".md",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".rtf",
    }
)

IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}


@dataclass
class IngestResult:
    document_text: str
    page_texts: Optional[list[str]]
    pdf_bytes: Optional[bytes]
    image_base64: Optional[str]
    image_media_type: Optional[str]
    source_format: str
    ingest_meta: dict


def supported_extensions_hint() -> str:
    exts = sorted(SUPPORTED_EXTENSIONS)
    return ", ".join(exts)


def ingest_upload(filename: str, contents: bytes) -> IngestResult:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        ext = _sniff_extension(contents) or ext

    if ext == ".pdf" or (not ext and contents[:4] == b"%PDF"):
        return _ingest_pdf(contents)

    if ext in IMAGE_MEDIA_TYPES:
        return _ingest_image(contents, ext)

    if ext in (".xlsx", ".xlsm"):
        return _ingest_xlsx(contents, ext)

    if ext == ".xls":
        return _ingest_xls(contents)

    if ext in (".csv", ".tsv"):
        return _ingest_delimited(contents, ext)

    if ext == ".docx":
        return _ingest_docx(contents)

    if ext == ".pptx":
        return _ingest_pptx(contents)

    if ext in (".txt", ".md", ".json", ".xml", ".html", ".htm", ".rtf"):
        return _ingest_plain_text(contents, ext)

    # ZIP-based Office formats without correct extension
    if contents[:2] == b"PK":
        for handler in (_try_xlsx, _try_docx, _try_pptx):
            result = handler(contents)
            if result:
                return result

    text = _decode_text_bytes(contents)
    if text.strip():
        return IngestResult(
            document_text=text,
            page_texts=None,
            pdf_bytes=None,
            image_base64=None,
            image_media_type=None,
            source_format=ext or "text",
            ingest_meta={
                "readMethod": "Raw text decode (fallback)",
                "characterCount": len(text),
            },
        )

    raise ValueError(
        f"Could not read this file. Supported types: {supported_extensions_hint()}"
    )


def _sniff_extension(contents: bytes) -> Optional[str]:
    if contents[:4] == b"%PDF":
        return ".pdf"
    if contents[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(contents)) as zf:
                names = zf.namelist()
                if any(n.startswith("xl/") for n in names):
                    return ".xlsx"
                if any(n.startswith("word/") for n in names):
                    return ".docx"
                if any(n.startswith("ppt/") for n in names):
                    return ".pptx"
        except zipfile.BadZipFile:
            pass
    try:
        import filetype

        kind = filetype.guess(contents)
        if kind:
            mime = kind.mime
            if mime == "application/pdf":
                return ".pdf"
            if mime.startswith("image/"):
                return "." + kind.extension
    except Exception:
        pass
    return None


def _ingest_pdf(contents: bytes) -> IngestResult:
    reader = pypdf.PdfReader(io.BytesIO(contents))
    page_texts = [page.extract_text() or "" for page in reader.pages]
    document_text = "\n\n".join(page_texts)
    pages_with_text = [i + 1 for i, t in enumerate(page_texts) if (t or "").strip()]
    return IngestResult(
        document_text=document_text,
        page_texts=page_texts,
        pdf_bytes=contents,
        image_base64=None,
        image_media_type=None,
        source_format=".pdf",
        ingest_meta={
            "readMethod": "PDF text (pypdf) + AI vision on rendered pages when applicable",
            "pageCount": len(page_texts),
            "pagesWithText": pages_with_text,
            "characterCount": len(document_text),
        },
    )


def _ingest_image(contents: bytes, ext: str) -> IngestResult:
    return IngestResult(
        document_text="",
        page_texts=None,
        pdf_bytes=None,
        image_base64=base64.b64encode(contents).decode("utf-8"),
        image_media_type=IMAGE_MEDIA_TYPES.get(ext, "image/jpeg"),
        source_format=ext,
        ingest_meta={
            "readMethod": "Image file — AI vision (no embedded text layer)",
            "pages": [1],
            "characterCount": 0,
        },
    )


def _ingest_xlsx(contents: bytes, ext: str) -> IngestResult:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"=== SHEET: {sheet_name} ===")
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_str(c) for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        parts.append("\n".join(rows))
    sheet_names = list(wb.sheetnames)
    wb.close()
    document_text = "\n\n".join(parts)
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=ext,
        ingest_meta={
            "readMethod": "Excel workbook (openpyxl) — all sheets as tables",
            "sheets": sheet_names,
            "characterCount": len(document_text),
        },
    )


def _try_xlsx(contents: bytes) -> Optional[IngestResult]:
    try:
        return _ingest_xlsx(contents, ".xlsx")
    except Exception:
        return None


def _ingest_xls(contents: bytes) -> IngestResult:
    import xlrd

    book = xlrd.open_workbook(file_contents=contents)
    parts: list[str] = []
    for sheet in book.sheets():
        parts.append(f"=== SHEET: {sheet.name} ===")
        rows: list[str] = []
        for rx in range(sheet.nrows):
            cells = [_cell_str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols)]
            if any(cells):
                rows.append("\t".join(cells))
        parts.append("\n".join(rows))
    sheet_names = [s.name for s in book.sheets()]
    document_text = "\n\n".join(parts)
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=".xls",
        ingest_meta={
            "readMethod": "Excel legacy .xls (xlrd)",
            "sheets": sheet_names,
            "characterCount": len(document_text),
        },
    )


def _ingest_delimited(contents: bytes, ext: str) -> IngestResult:
    text = _decode_text_bytes(contents)
    delimiter = "\t" if ext == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    lines = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
    document_text = "\n".join(lines)
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=ext,
        ingest_meta={
            "readMethod": f"Delimited table ({ext})",
            "characterCount": len(document_text),
        },
    )


def _ingest_docx(contents: bytes) -> IngestResult:
    from docx import Document

    doc = Document(io.BytesIO(contents))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        parts.append("--- TABLE ---")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    document_text = "\n\n".join(parts)
    table_count = len(doc.tables)
    return _try_docx_result(document_text, table_count=table_count)


def _try_docx(contents: bytes) -> Optional[IngestResult]:
    try:
        return _ingest_docx(contents)
    except Exception:
        return None


def _try_docx_result(document_text: str, table_count: int = 0) -> IngestResult:
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=".docx",
        ingest_meta={
            "readMethod": "Word document (python-docx) — paragraphs + tables",
            "tableCount": table_count,
            "characterCount": len(document_text),
        },
    )


def _ingest_pptx(contents: bytes) -> IngestResult:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(contents))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"=== SLIDE {i} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
    document_text = "\n\n".join(parts)
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=".pptx",
        ingest_meta={
            "readMethod": "PowerPoint (python-pptx) — slide text + tables",
            "slideCount": len(prs.slides),
            "characterCount": len(document_text),
        },
    )


def _try_pptx(contents: bytes) -> Optional[IngestResult]:
    try:
        return _ingest_pptx(contents)
    except Exception:
        return None


def _ingest_plain_text(contents: bytes, ext: str) -> IngestResult:
    document_text = _decode_text_bytes(contents)
    if ext == ".json":
        try:
            parsed = json.loads(document_text)
            document_text = json.dumps(parsed, indent=2)
        except json.JSONDecodeError:
            pass
    return IngestResult(
        document_text=document_text,
        page_texts=[document_text],
        pdf_bytes=None,
        image_base64=None,
        image_media_type=None,
        source_format=ext,
        ingest_meta={
            "readMethod": f"Plain text / markup ({ext})",
            "characterCount": len(document_text),
        },
    )


def _decode_text_bytes(contents: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return contents.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        import chardet

        detected = chardet.detect(contents)
        enc = detected.get("encoding") or "utf-8"
        return contents.decode(enc, errors="replace")
    except Exception:
        return contents.decode("utf-8", errors="replace")


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()
